"""Product Import R1: preview first, then one atomic create-only commit."""
from __future__ import annotations

import io
import uuid

from openpyxl import Workbook, load_workbook

from conftest import (
    auth,
    create_category,
    create_product,
    create_shop,
    new_seller,
    new_staff,
    seller_with_shop,
)
from fselling import models


HEADERS = [
    "Tên sản phẩm",
    "Mã sản phẩm",
    "Mã vạch",
    "Giá bán",
    "Giá vốn",
    "Tồn kho",
    "Danh mục",
]


def _csv(rows, delimiter=",") -> bytes:
    lines = [delimiter.join(HEADERS)]
    for row in rows:
        lines.append(delimiter.join(str(value) for value in row))
    return ("\ufeff" + "\n".join(lines)).encode("utf-8")


def _xlsx(rows, *, formula=False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sản phẩm"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    if formula:
        ws["D2"] = "=10000+5000"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _preview(client, token, shop_id, content, filename="san-pham.csv"):
    return client.post(
        f"/api/products/{shop_id}/imports/preview",
        files={"file": (filename, content, "application/octet-stream")},
        headers=auth(token),
    )


def _commit(client, token, shop_id, content, operation_id, filename="san-pham.csv"):
    return client.post(
        f"/api/products/{shop_id}/imports/commit",
        data={"operation_id": operation_id},
        files={"file": (filename, content, "application/octet-stream")},
        headers=auth(token),
    )


def test_preview_csv_utf8_semicolon_is_read_only(client, db):
    _, token = new_seller(client)
    shop_id = create_shop(client, token)
    content = _csv(
        [
            ["Cà phê sữa", "CF-01", "8934567890123", 25000, 15000, 12, "Đồ uống"],
            ["Bánh quy", "BQ-01", "", 32000, "", 8, "Bánh kẹo"],
        ],
        delimiter=";",
    )

    before_categories = db.query(models.Category).filter_by(shop_id=shop_id).count()
    before_products = db.query(models.Product).filter_by(shop_id=shop_id).count()
    response = _preview(client, token, shop_id, content)

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["can_commit"] is True
    assert body["stats"] == {"create": 2, "skip": 0, "error": 0}
    assert body["missing_categories"] == ["Bánh kẹo", "Đồ uống"]
    assert [row["action"] for row in body["rows"]] == ["CREATE", "CREATE"]
    assert db.query(models.Category).filter_by(shop_id=shop_id).count() == before_categories
    assert db.query(models.Product).filter_by(shop_id=shop_id).count() == before_products


def test_preview_xlsx_rejects_formula_and_duplicate_rows(client):
    _, token = new_seller(client)
    shop_id = create_shop(client, token)
    content = _xlsx(
        [
            ["Mì gói", "MI-01", "8934000000001", 5000, 3500, 20, "Đồ khô"],
            ["Mì gói", "MI-01", "8934000000001", 5000, 3500, 20, "Đồ khô"],
        ],
        formula=True,
    )

    response = _preview(client, token, shop_id, content, "san-pham.xlsx")

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["can_commit"] is False
    assert body["stats"]["error"] == 2
    assert any("công thức" in error.lower() for error in body["rows"][0]["errors"])
    assert any("trùng" in error.lower() for error in body["rows"][1]["errors"])


def test_existing_product_is_skipped_not_updated(client, db):
    ctx = seller_with_shop(client)
    existing = client.post(
        "/api/products",
        params={"shop_id": ctx["shop_id"]},
        data={
            "name": "Nước suối",
            "code": "NS-01",
            "barcode": "8934000000099",
            "price": 5000,
            "stock": 10,
            "category_id": ctx["category_id"],
        },
        headers=auth(ctx["token"]),
    )
    assert existing.status_code == 200, existing.text
    content = _csv(
        [["Tên mới không được ghi", "NS-01", "", 99999, 1, 999, "Danh mục mới"]]
    )

    preview = _preview(client, ctx["token"], ctx["shop_id"], content).json()
    assert preview["stats"] == {"create": 0, "skip": 1, "error": 0}
    assert preview["rows"][0]["action"] == "SKIP"

    operation_id = uuid.uuid4().hex
    response = _commit(client, ctx["token"], ctx["shop_id"], content, operation_id)
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 0
    db.expire_all()
    product = db.get(models.Product, existing.json()["id"])
    assert product.name == "Nước suối"
    assert product.price == 5000
    assert product.stock == 10


def test_commit_is_atomic_idempotent_and_undo_hides_batch(client, db):
    _, token = new_seller(client)
    shop_id = create_shop(client, token)
    content = _xlsx(
        [
            ["Sữa hộp", "SH-01", "8934000000101", 12000, 8000, 24, "Sữa"],
            ["Khăn giấy", "KG-01", "", 18000, "", 15, "Gia dụng"],
        ]
    )
    operation_id = uuid.uuid4().hex

    response = _commit(
        client, token, shop_id, content, operation_id, "san-pham.xlsx"
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["created"] == 2
    assert body["skipped"] == 0
    assert body["categories_created"] == 2
    assert body["can_undo"] is True
    product_ids = body["product_ids"]

    replay = _commit(client, token, shop_id, content, operation_id, "san-pham.xlsx")
    assert replay.status_code == 200, replay.text
    assert replay.json()["replayed"] is True
    assert replay.json()["product_ids"] == product_ids
    assert db.query(models.Product).filter_by(shop_id=shop_id).count() == 2

    undo = client.post(
        f"/api/products/{shop_id}/imports/{operation_id}/undo",
        headers=auth(token),
    )
    assert undo.status_code == 200, undo.text
    assert undo.json() == {"hidden": 2, "already_undone": False}
    db.expire_all()
    assert all(db.get(models.Product, product_id).is_active is False for product_id in product_ids)

    undo_again = client.post(
        f"/api/products/{shop_id}/imports/{operation_id}/undo",
        headers=auth(token),
    )
    assert undo_again.status_code == 200
    assert undo_again.json() == {"hidden": 0, "already_undone": True}


def test_invalid_commit_writes_nothing(client, db):
    _, token = new_seller(client)
    shop_id = create_shop(client, token)
    content = _csv(
        [
            ["Hợp lệ", "OK-01", "", 10000, 6000, 5, "Mới"],
            ["Giá sai", "BAD-01", "", -1, 0, 5, "Mới"],
        ]
    )

    response = _commit(client, token, shop_id, content, uuid.uuid4().hex)

    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "PRODUCT_IMPORT_INVALID"
    assert db.query(models.Product).filter_by(shop_id=shop_id).count() == 0
    assert db.query(models.Category).filter_by(shop_id=shop_id).count() == 0


def test_import_is_owner_only_and_tenant_bound(client):
    owner = seller_with_shop(client)
    _, staff_token = new_staff(client, owner, "WAREHOUSE")
    _, other_token = new_seller(client)
    content = _csv([["Kẹo", "KEO-01", "", 1000, 500, 10, "Bánh kẹo"]])

    staff = _preview(client, staff_token, owner["shop_id"], content)
    other = _preview(client, other_token, owner["shop_id"], content)

    assert staff.status_code == 404
    assert other.status_code == 404


def test_file_boundary_and_required_headers(client):
    _, token = new_seller(client)
    shop_id = create_shop(client, token)

    wrong_type = _preview(client, token, shop_id, b"hello", "products.txt")
    missing_headers = _preview(
        client, token, shop_id, "Tên sản phẩm,Giá bán\nKẹo,1000".encode("utf-8")
    )
    too_large = _preview(client, token, shop_id, b"x" * (2 * 1024 * 1024 + 1))

    assert wrong_type.status_code == 400
    assert missing_headers.status_code == 400
    assert too_large.status_code == 413


def test_libreoffice_template_has_exact_import_headers():
    path = (
        __import__("pathlib").Path(__file__).resolve().parent.parent
        / "static"
        / "templates"
        / "fselling-mau-nhap-san-pham.xlsx"
    )
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet = workbook["Sản phẩm"]
    assert [sheet.cell(1, col).value for col in range(1, 8)] == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet.max_row >= 4
    workbook.close()
