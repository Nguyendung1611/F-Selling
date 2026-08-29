"""Owner-only, preview-first product import for LibreOffice XLSX/CSV files."""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any, Iterable

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .. import models
from ..core.i18n import tr
from ..core.money import checked_multiply
from ..core.numeric_limits import MAX_SAFE_QUANTITY, MAX_SAFE_VND
from ..dependencies import require_own_shop
from . import catalog_service, inventory_service


MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_XLSX_BYTES = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
_OPERATION_ID = re.compile(r"^[A-Za-z0-9_-]{8,64}$")

_HEADERS = {
    "name": {"ten_san_pham", "product_name", "name"},
    "code": {"ma_san_pham", "sku", "code"},
    "barcode": {"ma_vach", "barcode"},
    "price": {"gia_ban", "sale_price", "price"},
    "cost_price": {"gia_von", "cost_price"},
    "stock": {"ton_kho", "opening_stock", "stock"},
    "category": {"danh_muc", "category"},
}
_REQUIRED_HEADERS = {"name", "price", "stock", "category"}
_HEADER_LABELS = {
    "name": "Tên sản phẩm",
    "code": "Mã sản phẩm",
    "barcode": "Mã vạch",
    "price": "Giá bán",
    "cost_price": "Giá vốn",
    "stock": "Tồn kho",
    "category": "Danh mục",
}


def _normalize_header(value: Any) -> str:
    plain = unicodedata.normalize("NFKD", str(value or ""))
    plain = "".join(ch for ch in plain if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "_", plain.casefold()).strip("_")


def _header_map(values: Iterable[Any]) -> dict[str, int]:
    aliases = {alias: key for key, group in _HEADERS.items() for alias in group}
    result: dict[str, int] = {}
    for index, value in enumerate(values):
        key = aliases.get(_normalize_header(value))
        if key and key not in result:
            result[key] = index
    missing = sorted(_REQUIRED_HEADERS - result.keys())
    if missing:
        labels = ", ".join(_HEADER_LABELS[key] for key in missing)
        raise HTTPException(
            status_code=400,
            detail=tr("File đang thiếu cột bắt buộc: {labels}", labels=labels),
        )
    return result


def _text(value: Any, *, label: str, maximum: int, required: bool = False) -> str:
    text = " ".join(str(value or "").split())
    if required and not text:
        raise ValueError(f"{label} không được để trống")
    if len(text) > maximum:
        raise ValueError(f"{label} tối đa {maximum} ký tự")
    if text.startswith(("=", "+", "@", "\t", "\r")):
        raise ValueError(f"{label} không được bắt đầu bằng ký hiệu công thức")
    return text


def _integer(value: Any, *, label: str, required: bool) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ValueError(f"{label} không được để trống")
        return None
    if isinstance(value, bool):
        raise ValueError(f"{label} phải là số nguyên")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"{label} phải là số nguyên")
        return int(value)
    raw = str(value).strip().replace("₫", "").replace("đ", "").replace(" ", "")
    if re.fullmatch(r"[+-]?\d+", raw):
        return int(raw)
    if re.fullmatch(r"[+-]?\d{1,3}(?:[.,]\d{3})+", raw):
        return int(raw.replace(".", "").replace(",", ""))
    raise ValueError(f"{label} phải là số nguyên; ví dụ 25000")


def _read_csv(content: bytes) -> tuple[list[Any], list[tuple[int, list[Any], set[int]]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=400,
            detail=tr("File CSV phải được lưu với mã hóa UTF-8"),
        ) from exc
    if "\x00" in text:
        raise HTTPException(status_code=400, detail=tr("File CSV không hợp lệ"))
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        raise HTTPException(status_code=400, detail=tr("File không có dòng tiêu đề"))
    return rows[0], [
        (index, row, set()) for index, row in enumerate(rows[1:], start=2)
    ]


def _read_xlsx(content: bytes) -> tuple[list[Any], list[tuple[int, list[Any], set[int]]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 200 or sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_XLSX_BYTES:
                raise HTTPException(
                    status_code=400,
                    detail=tr("File XLSX chứa quá nhiều dữ liệu để nhập an toàn"),
                )
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=400, detail=tr("File XLSX không hợp lệ")) from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=tr("Không đọc được file XLSX")) from exc
    try:
        worksheet = workbook["Sản phẩm"] if "Sản phẩm" in workbook.sheetnames else workbook.worksheets[0]
        iterator = worksheet.iter_rows()
        header_cells = next(iterator, None)
        if not header_cells:
            raise HTTPException(status_code=400, detail=tr("File không có dòng tiêu đề"))
        header = [cell.value for cell in header_cells]
        rows: list[tuple[int, list[Any], set[int]]] = []
        for row_number, cells in enumerate(iterator, start=2):
            rows.append(
                (
                    row_number,
                    [cell.value for cell in cells],
                    {index for index, cell in enumerate(cells) if cell.data_type == "f"},
                )
            )
        return header, rows
    finally:
        workbook.close()


def _read_rows(filename: str, content: bytes):
    if not content:
        raise HTTPException(status_code=400, detail=tr("File đang rỗng"))
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail=tr("File quá lớn; tối đa 2 MB"))
    extension = Path(filename or "").suffix.casefold()
    if extension == ".csv":
        header, rows = _read_csv(content)
    elif extension == ".xlsx":
        header, rows = _read_xlsx(content)
    else:
        raise HTTPException(
            status_code=400,
            detail=tr("Chỉ chấp nhận file .xlsx hoặc .csv"),
        )
    mapping = _header_map(header)
    populated = [
        row for row in rows if any(value not in (None, "") for value in row[1])
    ]
    if len(populated) > MAX_IMPORT_ROWS:
        raise HTTPException(
            status_code=400,
            detail=tr("File có quá 1.000 dòng sản phẩm"),
        )
    if not populated:
        raise HTTPException(status_code=400, detail=tr("File chưa có sản phẩm nào"))
    return mapping, populated


def _cell(values: list[Any], mapping: dict[str, int], key: str) -> Any:
    index = mapping.get(key)
    return values[index] if index is not None and index < len(values) else None


def _existing_keys(db: Session, shop_id: int):
    products = db.query(models.Product).filter(models.Product.shop_id == shop_id).all()
    return (
        {" ".join((p.name or "").split()).casefold(): p for p in products},
        {(p.code or "").strip().casefold(): p for p in products if (p.code or "").strip()},
        {(p.barcode or "").strip().upper(): p for p in products if (p.barcode or "").strip()},
    )


def preview_product_import(
    db: Session,
    current_user: models.User,
    shop_id: int,
    filename: str,
    content: bytes,
) -> dict:
    require_own_shop(db, shop_id, current_user)
    mapping, source_rows = _read_rows(filename, content)
    existing_names, existing_codes, existing_barcodes = _existing_keys(db, shop_id)
    categories = db.query(models.Category).filter(models.Category.shop_id == shop_id).all()
    category_by_name = {" ".join(c.name.split()).casefold(): c for c in categories}

    seen_names: set[str] = set()
    seen_codes: set[str] = set()
    seen_barcodes: set[str] = set()
    result_rows: list[dict] = []
    missing_categories: set[str] = set()
    stats = {"create": 0, "skip": 0, "error": 0}

    for row_number, values, formula_columns in source_rows:
        errors: list[str] = []
        data: dict[str, Any] = {}
        for key, index in mapping.items():
            if index in formula_columns:
                errors.append(
                    f"Không dùng công thức ở cột {_HEADER_LABELS[key]}; hãy dán giá trị"
                )
        try:
            data["name"] = _text(
                _cell(values, mapping, "name"),
                label="Tên sản phẩm",
                maximum=300,
                required=True,
            )
        except ValueError as exc:
            errors.append(str(exc))
            data["name"] = ""
        for key, label in (("code", "Mã sản phẩm"), ("barcode", "Mã vạch")):
            try:
                data[key] = _text(
                    _cell(values, mapping, key), label=label, maximum=64
                )
            except ValueError as exc:
                errors.append(str(exc))
                data[key] = ""
        try:
            data["barcode"] = catalog_service.normalize_barcode(data["barcode"])
        except HTTPException as exc:
            errors.append(str(exc.detail))
            data["barcode"] = None
        try:
            data["category"] = _text(
                _cell(values, mapping, "category"),
                label="Danh mục",
                maximum=120,
                required=True,
            )
        except ValueError as exc:
            errors.append(str(exc))
            data["category"] = ""

        for key, label, required in (
            ("price", "Giá bán", True),
            ("cost_price", "Giá vốn", False),
            ("stock", "Tồn kho", True),
        ):
            try:
                data[key] = _integer(
                    _cell(values, mapping, key), label=label, required=required
                )
            except ValueError as exc:
                errors.append(str(exc))
                data[key] = None

        if data.get("price") is not None and not 0 < data["price"] <= MAX_SAFE_VND:
            errors.append("Giá bán phải lớn hơn 0 và nằm trong giới hạn VND")
        if data.get("cost_price") is not None and not 0 <= data["cost_price"] <= MAX_SAFE_VND:
            errors.append("Giá vốn phải từ 0 trở lên và nằm trong giới hạn VND")
        if data.get("stock") is not None and not 0 <= data["stock"] <= MAX_SAFE_QUANTITY:
            errors.append("Tồn kho phải từ 0 trở lên và nằm trong giới hạn")
        if (
            data.get("stock") is not None
            and data.get("cost_price") is not None
            and data["stock"] * data["cost_price"] > MAX_SAFE_VND
        ):
            errors.append("Tổng giá vốn tồn đầu vượt giới hạn VND")

        name_key = data.get("name", "").casefold()
        code_key = data.get("code", "").casefold()
        barcode_key = (data.get("barcode") or "").upper()
        duplicate_fields = []
        if name_key and name_key in seen_names:
            duplicate_fields.append("tên sản phẩm")
        if code_key and code_key in seen_codes:
            duplicate_fields.append("mã sản phẩm")
        if barcode_key and barcode_key in seen_barcodes:
            duplicate_fields.append("mã vạch")
        if duplicate_fields:
            errors.append("Dòng bị trùng " + ", ".join(duplicate_fields) + " trong file")
        if name_key:
            seen_names.add(name_key)
        if code_key:
            seen_codes.add(code_key)
        if barcode_key:
            seen_barcodes.add(barcode_key)

        matched = (
            existing_names.get(name_key)
            or (existing_codes.get(code_key) if code_key else None)
            or (existing_barcodes.get(barcode_key) if barcode_key else None)
        )
        category = category_by_name.get(data.get("category", "").casefold())
        if category is not None and category.is_active is False:
            errors.append("Danh mục đang bị ẩn; hãy bật lại trước khi nhập")

        if errors:
            action = "ERROR"
            stats["error"] += 1
        elif matched is not None:
            action = "SKIP"
            stats["skip"] += 1
        else:
            action = "CREATE"
            stats["create"] += 1
            if category is None:
                missing_categories.add(data["category"])
        result_rows.append(
            {
                "row_number": row_number,
                "action": action,
                "data": data,
                "errors": errors,
                "matched_product_id": matched.id if matched is not None else None,
            }
        )

    return {
        "filename": Path(filename or "").name[:120],
        "total_rows": len(result_rows),
        "stats": stats,
        "missing_categories": sorted(missing_categories, key=str.casefold),
        "can_commit": stats["error"] == 0 and bool(result_rows),
        "rows": result_rows,
        "source_sha256": hashlib.sha256(content).hexdigest(),
    }


def _validate_operation_id(operation_id: str) -> str:
    value = (operation_id or "").strip()
    if not _OPERATION_ID.fullmatch(value):
        raise HTTPException(status_code=400, detail=tr("Mã lần nhập không hợp lệ"))
    return value


def _find_import_log(db: Session, shop_id: int, operation_id: str, action: str):
    marker = f"import_id={operation_id};"
    return (
        db.query(models.SystemLog)
        .filter(
            models.SystemLog.shop_id == shop_id,
            models.SystemLog.action == action,
            models.SystemLog.details.contains(marker),
        )
        .order_by(models.SystemLog.id.desc())
        .first()
    )


def _detail_value(details: str, key: str) -> str:
    match = re.search(rf"(?:^|;){re.escape(key)}=([^;]*);", details or "")
    return match.group(1) if match else ""


def _detail_ids(details: str, key: str) -> list[int]:
    raw = _detail_value(details, key)
    return [int(value) for value in raw.split(",") if value.isdigit()]


def _replay_response(log: models.SystemLog, source_sha256: str) -> dict:
    details = log.details or ""
    if _detail_value(details, "sha256") != source_sha256:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "PRODUCT_IMPORT_OPERATION_CONFLICT",
                "message": tr("Mã lần nhập đã được dùng cho một file khác"),
            },
        )
    product_ids = _detail_ids(details, "product_ids")
    return {
        "operation_id": _detail_value(details, "import_id"),
        "created": int(_detail_value(details, "created") or 0),
        "skipped": int(_detail_value(details, "skipped") or 0),
        "categories_created": int(_detail_value(details, "categories_created") or 0),
        "product_ids": product_ids,
        "can_undo": bool(product_ids),
        "replayed": True,
    }


def commit_product_import(
    db: Session,
    current_user: models.User,
    shop_id: int,
    filename: str,
    content: bytes,
    operation_id: str,
) -> dict:
    require_own_shop(db, shop_id, current_user)
    operation_id = _validate_operation_id(operation_id)
    source_sha256 = hashlib.sha256(content).hexdigest()
    inventory_service.lock_shop_for_inventory(db, shop_id)
    db.expire_all()

    replay = _find_import_log(db, shop_id, operation_id, "IMPORT_PRODUCTS")
    if replay is not None:
        db.rollback()
        return _replay_response(replay, source_sha256)

    preview = preview_product_import(
        db, current_user, shop_id, filename, content
    )
    if not preview["can_commit"]:
        db.rollback()
        raise HTTPException(
            status_code=400,
            detail={
                "code": "PRODUCT_IMPORT_INVALID",
                "message": tr("File còn dòng lỗi; chưa có dữ liệu nào được nhập"),
                "preview": preview,
            },
        )

    category_by_name = {
        " ".join(category.name.split()).casefold(): category
        for category in db.query(models.Category).filter(models.Category.shop_id == shop_id)
    }
    created_categories: list[models.Category] = []
    for name in preview["missing_categories"]:
        category = models.Category(name=name, shop_id=shop_id, is_active=True)
        db.add(category)
        created_categories.append(category)
        category_by_name[name.casefold()] = category
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail=tr("Dữ liệu kho vừa thay đổi; hãy xem trước lại file"),
        ) from exc

    products: list[models.Product] = []
    for row in preview["rows"]:
        if row["action"] != "CREATE":
            continue
        data = row["data"]
        category = category_by_name[data["category"].casefold()]
        cost_price = data["cost_price"]
        stock = data["stock"]
        product = models.Product(
            code=data["code"] or None,
            barcode=data["barcode"],
            name=data["name"],
            price=data["price"],
            stock=stock,
            cost_known_qty=stock if cost_price is not None else 0,
            cost_unknown_qty=stock if cost_price is None else 0,
            cost_basis_vnd=checked_multiply(stock, cost_price or 0),
            cost_deficit_qty=0,
            cost_state_version=1 if stock else 0,
            image_url=catalog_service.DEFAULT_PRODUCT_IMAGE,
            category_id=category.id,
            shop_id=shop_id,
            track_batches=False,
        )
        db.add(product)
        products.append(product)
    try:
        db.flush()
        for product in products:
            if not product.code:
                product.code = f"SP-{product.id}"
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail=tr("Dữ liệu kho vừa thay đổi; hãy xem trước lại file"),
        ) from exc

    product_ids = [int(product.id) for product in products]
    category_ids = [int(category.id) for category in created_categories]
    details = (
        f"import_id={operation_id};sha256={source_sha256};"
        f"created={len(product_ids)};skipped={preview['stats']['skip']};"
        f"categories_created={len(category_ids)};"
        f"product_ids={','.join(map(str, product_ids))};"
        f"category_ids={','.join(map(str, category_ids))};"
    )
    db.add(
        models.SystemLog(
            user_id=current_user.id,
            shop_id=shop_id,
            action="IMPORT_PRODUCTS",
            details=details,
        )
    )
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail=tr("Dữ liệu kho vừa thay đổi; hãy xem trước lại file"),
        ) from exc
    return {
        "operation_id": operation_id,
        "created": len(product_ids),
        "skipped": preview["stats"]["skip"],
        "categories_created": len(category_ids),
        "product_ids": product_ids,
        "can_undo": bool(product_ids),
        "replayed": False,
    }


def undo_product_import(
    db: Session,
    current_user: models.User,
    shop_id: int,
    operation_id: str,
) -> dict:
    require_own_shop(db, shop_id, current_user)
    operation_id = _validate_operation_id(operation_id)
    inventory_service.lock_shop_for_inventory(db, shop_id)
    db.expire_all()
    source = _find_import_log(db, shop_id, operation_id, "IMPORT_PRODUCTS")
    if source is None:
        db.rollback()
        raise HTTPException(status_code=404, detail=tr("Không tìm thấy lần nhập này"))
    if _find_import_log(db, shop_id, operation_id, "UNDO_IMPORT_PRODUCTS") is not None:
        db.rollback()
        return {"hidden": 0, "already_undone": True}

    product_ids = _detail_ids(source.details or "", "product_ids")
    products = (
        db.query(models.Product)
        .filter(
            models.Product.shop_id == shop_id,
            models.Product.id.in_(product_ids or [-1]),
            models.Product.is_active == True,  # noqa: E712
        )
        .all()
    )
    for product in products:
        product.is_active = False
    db.add(
        models.SystemLog(
            user_id=current_user.id,
            shop_id=shop_id,
            action="UNDO_IMPORT_PRODUCTS",
            details=f"import_id={operation_id};hidden={len(products)};",
        )
    )
    db.commit()
    return {"hidden": len(products), "already_undone": False}


__all__ = [
    "MAX_IMPORT_BYTES",
    "commit_product_import",
    "preview_product_import",
    "undo_product_import",
]
